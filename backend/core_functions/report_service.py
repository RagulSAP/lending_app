"""
Report data queries and export helpers (Excel via openpyxl, PDF via reportlab).
"""
import io
from datetime import datetime

from sqlalchemy import func

from models import Transaction, Loan, Expense, Customer, User, ExpenseCategory


# ---------------------------------------------------------------------------
# Column definitions (used for both Excel and PDF)
# ---------------------------------------------------------------------------

REPORT_HEADERS = {
    "transactions": ["Date", "Customer", "Amount", "Mode", "Type", "Collected By", "Remarks"],
    "loans": ["Loan ID", "Customer", "Disbursed", "Rate %", "Interest Type",
              "Installment Type", "Total Payable", "Total Paid", "Balance", "Status"],
    "expenses": ["Date", "Category", "Amount", "Entered By", "Remark"],
}


def _txn_row(r):
    return [
        r.get("transaction_date", ""),
        r.get("customer_name", ""),
        r.get("amount", ""),
        r.get("payment_mode", ""),
        r.get("transaction_type", ""),
        r.get("collected_by", ""),
        r.get("remarks", ""),
    ]


def _loan_row(r):
    return [
        r.get("loan_id", ""),
        r.get("customer_name", ""),
        r.get("disbursement_amount", ""),
        r.get("interest_rate", ""),
        r.get("interest_type", ""),
        r.get("installment_type", ""),
        r.get("total_payable", ""),
        r.get("total_paid", ""),
        r.get("balance_amount", ""),
        r.get("status", ""),
    ]


def _expense_row(r):
    return [
        r.get("expense_date", ""),
        r.get("category_name", ""),
        r.get("expense_amount", ""),
        r.get("user_name", ""),
        r.get("expense_remark", ""),
    ]


_ROW_BUILDERS = {
    "transactions": _txn_row,
    "loans": _loan_row,
    "expenses": _expense_row,
}


# ---------------------------------------------------------------------------
# Data query functions
# ---------------------------------------------------------------------------

def get_transactions_report(
    db,
    org_id: str,
    date_from,
    date_to,
    payment_mode=None,
    transaction_type=None,
    customer_id=None,
    page: int = 1,
    per_page: int = 50,
):
    """
    Returns (records_list, total_count) for transactions within the given org and date range.
    Joins Customer and User to provide human-readable names.
    """
    q = (
        db.query(
            Transaction,
            Customer.name.label("customer_name"),
            User.name.label("user_name"),
        )
        .join(Customer, Customer.customer_id == Transaction.customer_id, isouter=True)
        .join(User, User.user_id == Transaction.user_id, isouter=True)
        .filter(
            Transaction.org_id == org_id,
            Transaction.transaction_date >= date_from,
            Transaction.transaction_date <= date_to,
        )
    )
    if payment_mode:
        q = q.filter(Transaction.payment_mode == payment_mode)
    if transaction_type:
        q = q.filter(Transaction.transaction_type == transaction_type)
    if customer_id:
        q = q.filter(Transaction.customer_id == customer_id)

    total = q.count()
    rows = (
        q.order_by(Transaction.transaction_date.desc(), Transaction.created_at.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )

    records = []
    for txn, customer_name, user_name in rows:
        records.append({
            "transaction_id": txn.transaction_id,
            "transaction_date": txn.transaction_date.isoformat() if txn.transaction_date else None,
            "customer_name": customer_name,
            "amount": float(txn.amount) if txn.amount is not None else 0.0,
            "payment_mode": txn.payment_mode,
            "transaction_type": txn.transaction_type,
            "collected_by": user_name,
            "remarks": txn.remarks,
            "loan_id": txn.loan_id,
            "installment_id": txn.installment_id,
        })
    return records, total


def get_loans_report(
    db,
    org_id: str,
    date_from,
    date_to,
    status=None,
    interest_type=None,
    page: int = 1,
    per_page: int = 50,
):
    """Returns (records_list, total_count) for loans matching the given filters."""
    q = (
        db.query(Loan, Customer.name.label("customer_name"))
        .join(Customer, Customer.customer_id == Loan.customer_id, isouter=True)
        .filter(Loan.org_id == org_id)
    )
    if date_from:
        q = q.filter(Loan.disbursement_date >= date_from)
    if date_to:
        q = q.filter(Loan.disbursement_date <= date_to)
    if status:
        q = q.filter(Loan.status == status)
    if interest_type:
        q = q.filter(Loan.interest_type == interest_type)

    total = q.count()
    rows = (
        q.order_by(Loan.disbursement_date.desc(), Loan.created_at.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )

    records = []
    for loan, customer_name in rows:
        records.append({
            "loan_id": loan.loan_id,
            "customer_name": customer_name,
            "disbursement_amount": float(loan.disbursement_amount) if loan.disbursement_amount else 0.0,
            "interest_rate": float(loan.interest_rate) if loan.interest_rate else 0.0,
            "interest_type": loan.interest_type,
            "installment_type": loan.installment_type,
            "total_payable": float(loan.total_payable) if loan.total_payable else 0.0,
            "total_paid": float(loan.total_paid) if loan.total_paid else 0.0,
            "balance_amount": float(loan.balance_amount) if loan.balance_amount else 0.0,
            "status": loan.status,
            "disbursement_date": loan.disbursement_date.isoformat() if loan.disbursement_date else None,
        })
    return records, total


def get_expenses_report(
    db,
    org_id: str,
    date_from,
    date_to,
    category_id=None,
    page: int = 1,
    per_page: int = 50,
):
    """Returns (records_list, total_count) for expenses within the org and date range."""
    q = (
        db.query(
            Expense,
            ExpenseCategory.name.label("category_name"),
            User.name.label("user_name"),
        )
        .join(ExpenseCategory, ExpenseCategory.category_id == Expense.category_id)
        .join(User, User.user_id == Expense.user_id, isouter=True)
        .filter(ExpenseCategory.org_id == org_id)
    )
    if date_from:
        q = q.filter(Expense.expense_date >= date_from)
    if date_to:
        q = q.filter(Expense.expense_date <= date_to)
    if category_id:
        q = q.filter(Expense.category_id == category_id)

    total = q.count()
    rows = (
        q.order_by(Expense.expense_date.desc(), Expense.created_at.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
        .all()
    )

    records = []
    for expense, category_name, user_name in rows:
        records.append({
            "expense_id": expense.expense_id,
            "expense_date": expense.expense_date.isoformat() if expense.expense_date else None,
            "category_name": category_name,
            "expense_amount": float(expense.expense_amount) if expense.expense_amount else 0.0,
            "user_name": user_name,
            "expense_remark": expense.expense_remark,
        })
    return records, total


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

_NAVY = "003366"
_WHITE = "FFFFFF"
_LIGHT_GREY = "F2F2F2"


def export_to_excel(report_type: str, records: list, filters_label: str = "") -> io.BytesIO:
    """
    Build an openpyxl workbook from report data.
    Returns a BytesIO buffer ready to send as a file download.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report_type.capitalize()

    header_font = Font(bold=True, color=_WHITE)
    header_fill = PatternFill("solid", fgColor=_NAVY)
    alt_fill = PatternFill("solid", fgColor=_LIGHT_GREY)
    center = Alignment(horizontal="center", vertical="center")

    # Title block
    ws.append([f"{report_type.upper()} REPORT"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=14)
    if filters_label:
        ws.append([filters_label])
    ws.append([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    ws.append([])

    headers = REPORT_HEADERS.get(report_type, [])
    ws.append(headers)
    header_row_idx = ws.max_row

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    row_builder = _ROW_BUILDERS.get(report_type, lambda r: list(r.values()))

    for data_row_offset, record in enumerate(records):
        row_data = row_builder(record)
        ws.append(row_data)
        actual_row = header_row_idx + 1 + data_row_offset
        if data_row_offset % 2 == 1:
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=actual_row, column=col_idx).fill = alt_fill

    # Auto-width
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(
            (len(str(cell.value)) if cell.value is not None else 0) for cell in col
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 55)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------

def export_to_pdf(report_type: str, records: list, filters_label: str = "") -> io.BytesIO:
    """
    Build a reportlab PDF from report data.
    Returns a BytesIO buffer ready to send as a file download.
    """
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"{report_type.upper()} REPORT", styles["Title"]))
    if filters_label:
        elements.append(Paragraph(filters_label, styles["Normal"]))
    elements.append(
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"])
    )
    elements.append(Spacer(1, 12))

    headers = REPORT_HEADERS.get(report_type, [])
    row_builder = _ROW_BUILDERS.get(report_type, lambda r: list(r.values()))

    table_data = [headers]
    for record in records:
        table_data.append([str(v) if v is not None else "" for v in row_builder(record)])

    col_count = len(headers)
    available_width = doc.width
    col_width = available_width / col_count

    t = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)

    row_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    t.setStyle(TableStyle(row_styles))
    elements.append(t)

    doc.build(elements)
    buf.seek(0)
    return buf
